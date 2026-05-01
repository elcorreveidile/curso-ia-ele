"""Course documents: DOCX sources stored under static_documents/, served as
DOCX or auto-rendered to a brand-consistent PDF using ReportLab.

Two documents are exposed:
  - bienvenida-videotutoria-1: welcome letter + Zoom invite for session 1
  - consentimiento-grabacion: GDPR recording consent (signature required)

Both documents are visible to enrolled students and admins, listed at
/api/documents and downloadable via /api/documents/{slug}/download.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


STATIC_DOCS = Path(__file__).parent / "static_documents"

# Brand palette (mirrors pdf_builder.py and the website).
NAVY = colors.HexColor("#0F4C81")
AMBER = colors.HexColor("#F5A623")
INK = colors.HexColor("#1A2535")
INK_SOFT = colors.HexColor("#46476A")
INK_MUTED = colors.HexColor("#6B82A0")


@dataclass(frozen=True)
class CourseDocument:
    slug: str
    title: str
    description: str
    docx_filename: str
    pdf_basename: str  # download filename for the rendered PDF
    icon: str          # emoji shown in the UI card
    requires_signature: bool

    @property
    def docx_path(self) -> Path:
        return STATIC_DOCS / self.docx_filename


COURSE_DOCUMENTS: list[CourseDocument] = [
    CourseDocument(
        slug="bienvenida-videotutoria-1",
        title="Bienvenida · Videotutoría 1",
        description=(
            "Carta de bienvenida con los datos de conexión a la primera "
            "videotutoría (Zoom · 4 mayo, 16:00) y la agenda de la sesión."
        ),
        docx_filename="Invitacion_Videotutoria1_Mayo4.docx",
        pdf_basename="Bienvenida-Videotutoria1.pdf",
        icon="👋",
        requires_signature=False,
    ),
    CourseDocument(
        slug="consentimiento-grabacion",
        title="Consentimiento de grabación · RGPD",
        description=(
            "Documento de consentimiento informado para la grabación de las "
            "videotutorías. Léelo, fírmalo y devuélvemelo antes del lunes."
        ),
        docx_filename="Consentimiento_Grabacion_IA_ELE.docx",
        pdf_basename="Consentimiento-Grabacion.pdf",
        icon="📝",
        requires_signature=True,
    ),
]


def get_document(slug: str) -> CourseDocument | None:
    return next((d for d in COURSE_DOCUMENTS if d.slug == slug), None)


# ───────────────────────── DOCX → text extraction ─────────────────────────

def _docx_paragraphs(path: Path) -> Iterable[tuple[str, str]]:
    """Yield (style, text) tuples preserving heading levels and bullet markers.

    ``style`` is one of: ``h1``, ``h2``, ``h3``, ``bullet``, ``body``.
    Empty paragraphs are skipped to keep PDF spacing tight.
    """
    doc = DocxDocument(str(path))
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_name = (p.style.name or "").lower() if p.style else ""
        if "heading 1" in style_name or "title" in style_name:
            yield "h1", text
        elif "heading 2" in style_name:
            yield "h2", text
        elif "heading 3" in style_name:
            yield "h3", text
        elif "list" in style_name or "bullet" in style_name:
            yield "bullet", text
        else:
            yield "body", text


# ──────────────────────────── PDF rendering ───────────────────────────────

def _build_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "h1": ParagraphStyle(
            "H1", parent=base["Heading1"], fontName="Helvetica-Bold",
            fontSize=20, leading=24, textColor=NAVY,
            spaceBefore=18, spaceAfter=12, alignment=TA_LEFT,
        ),
        "h2": ParagraphStyle(
            "H2", parent=base["Heading2"], fontName="Helvetica-Bold",
            fontSize=14, leading=18, textColor=NAVY,
            spaceBefore=14, spaceAfter=8, alignment=TA_LEFT,
        ),
        "h3": ParagraphStyle(
            "H3", parent=base["Heading3"], fontName="Helvetica-Bold",
            fontSize=12, leading=15, textColor=INK,
            spaceBefore=10, spaceAfter=6, alignment=TA_LEFT,
        ),
        "body": ParagraphStyle(
            "Body", parent=base["BodyText"], fontName="Helvetica",
            fontSize=10.5, leading=15, textColor=INK_SOFT,
            spaceAfter=6, alignment=TA_LEFT,
        ),
        "bullet": ParagraphStyle(
            "Bullet", parent=base["BodyText"], fontName="Helvetica",
            fontSize=10.5, leading=15, textColor=INK_SOFT,
            spaceAfter=4, leftIndent=14, bulletIndent=2,
            alignment=TA_LEFT,
        ),
        "label": ParagraphStyle(
            "Label", fontName="Helvetica-Bold", fontSize=9, leading=11,
            textColor=AMBER, alignment=TA_CENTER, spaceAfter=2,
        ),
        "brand": ParagraphStyle(
            "Brand", fontName="Helvetica-Bold", fontSize=22, leading=24,
            textColor=AMBER, alignment=TA_CENTER, spaceAfter=14,
        ),
        "footer": ParagraphStyle(
            "Footer", fontName="Helvetica", fontSize=8, leading=10,
            textColor=INK_MUTED, alignment=TA_CENTER,
        ),
    }


_INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_INLINE_LINK_RE = re.compile(r"(https?://[\w\-./?=&%#:+]+)")


def _format_inline(text: str) -> str:
    """Convert a tiny subset of Markdown (bold, raw URLs) to ReportLab markup."""
    out = _INLINE_BOLD_RE.sub(r"<b>\1</b>", text)
    out = _INLINE_LINK_RE.sub(r'<link href="\1" color="#0F4C81"><u>\1</u></link>', out)
    return out


def _draw_page_chrome(canvas, doc, title: str) -> None:
    """Header band + footer with page number on every page."""
    canvas.saveState()
    width, height = A4
    # Top amber accent bar
    canvas.setFillColor(AMBER)
    canvas.rect(0, height - 0.4 * cm, width, 0.4 * cm, stroke=0, fill=1)
    # Brand on top-left
    canvas.setFillColor(NAVY)
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawString(2 * cm, height - 1.0 * cm, "LA CLASE DIGITAL")
    canvas.setFillColor(INK_MUTED)
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(width - 2 * cm, height - 1.0 * cm, title)
    # Footer
    canvas.setStrokeColor(colors.HexColor("#E8EEF5"))
    canvas.setLineWidth(0.4)
    canvas.line(2 * cm, 1.6 * cm, width - 2 * cm, 1.6 * cm)
    canvas.setFillColor(INK_MUTED)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(2 * cm, 1.1 * cm, "laclasedigital.com · benitezl@go.ugr.es")
    canvas.drawRightString(width - 2 * cm, 1.1 * cm, f"Página {doc.page}")
    canvas.restoreState()


def render_document_pdf(slug: str) -> bytes:
    """Render the given course document to a brand-consistent PDF (in-memory)."""
    doc_meta = get_document(slug)
    if doc_meta is None:
        raise ValueError(f"Unknown document slug: {slug}")
    if not doc_meta.docx_path.exists():
        raise FileNotFoundError(f"DOCX missing: {doc_meta.docx_path}")

    styles = _build_styles()
    buf = io.BytesIO()
    pdf = BaseDocTemplate(
        buf, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.8 * cm, bottomMargin=2 * cm,
        title=doc_meta.title, author="La Clase Digital",
    )
    frame = Frame(
        pdf.leftMargin, pdf.bottomMargin,
        pdf.width, pdf.height,
        leftPadding=0, rightPadding=0,
        topPadding=0, bottomPadding=0,
        showBoundary=0,
    )
    pdf.addPageTemplates([
        PageTemplate(
            id="default", frames=[frame],
            onPage=lambda c, d: _draw_page_chrome(c, d, doc_meta.title),
        ),
    ])

    story = []
    # Header block
    story.append(Paragraph("[ | ]", styles["brand"]))
    story.append(Paragraph("LA CLASE DIGITAL", styles["label"]))
    story.append(Spacer(1, 0.4 * cm))
    story.append(Paragraph(doc_meta.title, styles["h1"]))
    sep = Table([[""]], colWidths=[16.5 * cm], rowHeights=[0.06 * cm])
    sep.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), AMBER)]))
    story.append(sep)
    story.append(Spacer(1, 0.4 * cm))

    # Body — DOCX paragraph stream, with simple inline formatting
    for kind, text in _docx_paragraphs(doc_meta.docx_path):
        # Skip the title/header paragraphs that duplicate what we just rendered.
        if kind in {"h1"} and text.lower().strip(": ·").startswith(
            ("documento de consentimiento", "asunto:", "la clase digital")
        ):
            continue
        formatted = _format_inline(text)
        if kind == "bullet":
            story.append(Paragraph(formatted, styles["bullet"], bulletText="•"))
        else:
            style = styles.get(kind, styles["body"])
            story.append(Paragraph(formatted, style))

    # Footer note for documents that need a signature
    if doc_meta.requires_signature:
        story.append(Spacer(1, 0.6 * cm))
        story.append(Paragraph(
            "<b>Cómo devolver este documento firmado:</b> imprime, firma a "
            "mano (o usa una herramienta de firma digital) y envíame el PDF "
            'firmado a <link href="mailto:benitezl@go.ugr.es" color="#0F4C81">'
            "<u>benitezl@go.ugr.es</u></link>.",
            styles["body"],
        ))

    pdf.build(story)
    return buf.getvalue()
